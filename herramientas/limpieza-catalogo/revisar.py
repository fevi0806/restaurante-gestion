import json, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
B = json.load(open(sys.argv[1]))['tables']
O = json.load(open('out/copia_limpia.json'))['tables']
D = json.load(open('decisiones.json'))
I = json.load(open('out/copia_limpia.informe.json'))
P0 = {p['id']: p for p in B['products']}; P1 = {p['id']: p for p in O['products']}
S0 = {s['id']: s for s in B['suppliers']}; S1 = {s['id']: s for s in O['suppliers']}
F1 = collections.defaultdict(list)
for f in O['product_formats']:
    if f['active']: F1[f['product_id']].append(f"{f['name']} ({f['factor']:g})")
op = {i['id'] for i in B['inventories'] if (i['notes'] or '').startswith('[arranque]')}
cnt = collections.defaultdict(float)
for l in B['inventory_lines']:
    if l['inventory_id'] in op: cnt[l['product_id']] += l['counted']

wb = Workbook(); F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A2D'); N = Font(name=F); Bf = Font(name=F, bold=True)
WARN = PatternFill('solid', fgColor='FFF2CC')
def sheet(title, headers, rows, widths, warn_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]: c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = N; c.alignment = Alignment(wrap_text=True, vertical='top')
        if warn_col is not None and row[warn_col].value: [setattr(c, 'fill', WARN) for c in row]
    for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws

def unitf(src, tgt, f):
    s, t = P0[src], P0[tgt]
    if f == 1 and s['unit'] == t['unit']: return ''
    return f"1 {s['unit']} de este = {f:g} {t['unit']}"
# Fusiones
rows = []
for m in D['merges']:
    t = m['target']; fin = P1.get(t, {}).get('name', P0[t]['name'])
    for s, f in m['sources']:
        rows.append([fin, P0[s]['name'], unitf(s, t, f), cnt.get(s) or '', m['note'] or ('mismo artículo con otro nombre' if m['name'] else 'línea de abono, regalo o rectificativa'), 'Sí' if m['review'] else ''])
rows.sort(key=lambda r: (r[5] != 'Sí', r[0].lower()))
sheet('Fusiones', ['Se queda como', 'Se une (desaparece)', 'Conversión', 'Contado en el inventario', 'Por qué', 'Revisar'], rows, [42, 52, 26, 14, 50, 9], warn_col=5)
# Bajas
rows = sorted([[P0[b['id']]['name'], P0[b['id']]['category'] or '', S0.get(P0[b['id']]['supplier_id'], {}).get('name', ''), b['why']] for b in D['bajas']], key=lambda r: (r[3], r[0].lower()))
sheet('Bajas', ['Artículo', 'Categoría que tenía', 'Proveedor', 'Por qué'], rows, [55, 22, 34, 50])
# Precios por caja
rows = []
for k, n in D['pack'].items():
    k = int(k); p0 = P0[k]; p1 = P1.get(k, {})
    rows.append([p1.get('name', p0['name']), p0['price'], n, p1.get('price'), ', '.join(F1.get(k, []))])
ws = sheet('Precios por caja', ['Artículo', 'Precio que tenía (€, era la caja)', 'Unidades por caja', 'Precio por unidad ahora (€)', 'Formatos'], rows, [50, 18, 12, 18, 30])
for r in range(2, ws.max_row + 1): ws.cell(r, 2).number_format = '#,##0.00'; ws.cell(r, 4).number_format = '#,##0.000'
# Cambios de nombre y categoría
rows = []
for p1 in O['products']:
    p0 = P0.get(p1['id'])
    if not p0 or not p1['active']: continue
    if p0['name'] != p1['name'] or (p0['category'] or '') != (p1['category'] or ''):
        rows.append([p0['name'], p1['name'] if p0['name'] != p1['name'] else '', p0['category'] or '(sin categoría)', p1['category']])
rows.sort(key=lambda r: (r[3], r[0].lower()))
sheet('Nombres y categorías', ['Artículo (antes)', 'Nombre nuevo', 'Categoría antes', 'Categoría ahora'], rows, [52, 42, 24, 28])
# Unidades
rows = []
for k, c in D.get('convert', {}).items():
    k = int(k); p0 = P0[k]; p1 = P1.get(k, {})
    rows.append([p1.get('name', p0['name']), p1.get('category', ''), 'ud', c['unit'], c['format'], round(c['factor'], 4), cnt.get(k) or ''])
rows.sort(key=lambda r: (r[1], r[0].lower()))
sheet('Unidades', ['Artículo', 'Categoría', 'Antes se controlaba en', 'Ahora en', 'Formato de compra / recuento', '1 envase = (kg o l)', 'Contado en el inventario (en envases)'], rows, [50, 26, 12, 10, 26, 14, 16])
# Sin precio
rows = sorted([[p['name'], p['category'], p['unit'], cnt.get(p['id']) or ''] for p in O['products'] if p['active'] and not p['price']], key=lambda r: (r[1], r[0].lower()))
sheet('Sin precio', ['Artículo', 'Categoría', 'Se controla en', 'Contado en el inventario'], rows, [44, 28, 14, 14])
# Catálogo final
rows = sorted([[p['category'], p['name'], p['unit'], p['price'], ', '.join(F1.get(p['id'], [])), S1.get(p['supplier_id'], {}).get('name', '')] for p in O['products'] if p['active']], key=lambda r: (r[0], r[1].lower()))
ws = sheet('Catálogo limpio', ['Categoría', 'Artículo', 'Se controla en', 'Precio (€ por kg / l / ud)', 'Formatos', 'Proveedor'], rows, [28, 52, 12, 14, 34, 36])
for r in range(2, ws.max_row + 1): ws.cell(r, 4).number_format = '#,##0.000'
# Proveedores
rows = []
for t, srcs in D['supplier_merges']:
    for s in srcs: rows.append([S0[s]['name'], 'Fusionado con ' + S1.get(t, S0[t])['name']])
for n in I['supBajas']: rows.append([n, 'Baja: no es proveedor de género (sin artículos ni facturas de compra)'])
sheet('Proveedores', ['Proveedor', 'Qué se ha hecho'], rows, [60, 70])
# Compras importadas
by = collections.defaultdict(lambda: [0, 0, 0.0])
for r in O['receipts']:
    if r.get('notes') == 'Importada de Excel':
        s = S1.get(r['supplier_id'], {}).get('name', '?'); by[s][0] += 1; by[s][2] += r['total']
for l in O['receipt_lines']:
    pass
rows = sorted([[k, v[0], round(v[2], 2)] for k, v in by.items()], key=lambda r: -r[2])
ws = sheet('Compras importadas', ['Proveedor', 'Facturas', 'Importe sin IVA (€)'], rows, [56, 10, 18])
n = ws.max_row
for r in range(2, n + 1): ws.cell(r, 3).number_format = '#,##0.00'
ws.cell(n + 1, 1, 'Total').font = Bf; ws.cell(n + 1, 2, f'=SUM(B2:B{n})').font = Bf; ws.cell(n + 1, 3, f'=SUM(C2:C{n})').font = Bf; ws.cell(n + 1, 3).number_format = '#,##0.00'

# Resumen
ws = wb['Sheet']; ws.title = 'Resumen'; wb.move_sheet('Resumen', -(len(wb.sheetnames) - 1))
L = [
 ('Limpieza del catálogo · El Mercado de Casas Viejas', None),
 (f"Hecha sobre la copia de seguridad del {json.load(open(sys.argv[1]))['created_at'][:10]}. Nada se ha cambiado todavía en la app: esto es la propuesta.", None),
 ('', None),
 ('Artículos antes', I['before']['products']), ('Artículos después', I['after']['products']),
 ('Fusionados (repetidos, abonos, regalos)', "=COUNTA(Fusiones!B:B)-1"), ('Dados de baja (no son género)', "=COUNTA(Bajas!A:A)-1"),
 ('Con el precio de caja corregido', "=COUNTA('Precios por caja'!A:A)-1"), ('Con nombre o categoría cambiados', "=COUNTA('Nombres y categorías'!A:A)-1"),
 ('Pasan de «unidades» a kg o litros (el envase queda como formato)', "=COUNTA(Unidades!A:A)-1"),
 ('Proveedores antes', I['before']['suppliers']), ('Proveedores después', I['after']['suppliers']),
 ('Facturas de compra importadas (mayo–septiembre)', I['imp']['invoices']), ('Líneas de compra importadas', I['imp']['lines']),
 ('', None),
 ('Qué revisar (en amarillo en «Fusiones»)', "=COUNTIF(Fusiones!F:F,\"Sí\")"),
 ('Artículos sin precio (casi todos dados de alta durante el inventario)', "=COUNTA('Sin precio'!A:A)-1"),
 ('', None),
 ('Lo contado en el inventario de arranque se conserva: si dos artículos se fusionan, sus recuentos se suman (con la conversión de unidades).', None),
 ('Todo lo que se usa en escandallos va en kg o litros; se sigue contando y comprando por envase (botella, garrafa, saco, lata…) porque el envase es un formato.', None),
 ('El stock actual no vale (se cargó con lo comprado). Queda bien al cerrar el inventario de arranque.', None),
 ('Las compras importadas sirven para el historial: precios, consumo y pedidos sugeridos. No cambian el stock final, porque el inventario de arranque lo deja como lo contado.', None),
]
for i, (a, b) in enumerate(L, 1):
    ws.cell(i, 1, a).font = Font(name=F, bold=(i == 1), size=14 if i == 1 else 11)
    if b is not None: ws.cell(i, 2, b).font = Bf
ws.column_dimensions['A'].width = 95; ws.column_dimensions['B'].width = 14
wb.save(sys.argv[2])
